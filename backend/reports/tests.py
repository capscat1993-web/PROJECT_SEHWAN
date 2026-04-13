from django.test import Client, TestCase

from openpyxl import load_workbook

from reports.services.health_export import export_health_excel
from reports.services.financial_health import calculate_health
from reports.views import _key_metrics_payload


class ApiSmokeTests(TestCase):
    databases = {"default"}

    def setUp(self):
        self.client = Client()

    def test_root_returns_platform_metadata(self):
        response = self.client.get("/")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["status"], "ok")

    def test_overview_endpoint_returns_summary(self):
        response = self.client.get("/api/overview")
        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertIn("total_companies", payload)
        self.assertIn("industry_count", payload)
        self.assertIn("top_industries", payload)

    def test_company_dashboard_returns_payload(self):
        response = self.client.get("/api/companies/400/dashboard")
        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertIn("company", payload)
        self.assertIn("tables", payload)
        self.assertIn("key_metrics", payload)
        self.assertTrue(payload["key_metrics"]["periods"])
        self.assertIn("주요재무지표", payload["tables"])
        self.assertEqual(payload["key_metrics"]["periods"], ["2019.12", "2020.12", "2021.12"])
        self.assertNotIn("손익계산서", payload["tables"])

    def test_key_metrics_normalizes_recent_money_values_to_million_won(self):
        payload = _key_metrics_payload(424)
        self.assertEqual(payload["unit"], "백만원")
        self.assertEqual(payload["metrics"]["매출액"]["2024.12"], 59357.0)
        self.assertEqual(payload["metrics"]["영업이익"]["2024.12"], 3005.0)
        self.assertEqual(payload["metrics"]["당기순이익"]["2024.12"], 2834.0)

    def test_health_export_uses_template_and_populates_current_company(self):
        buffer, company_name = export_health_excel(424)
        workbook = load_workbook(buffer, data_only=False)

        self.assertEqual(company_name, "(주)동남기계")
        self.assertEqual(
            workbook.sheetnames,
            ["재무데이터 입력", "재무비율 분석", "종합 평가표", "고객사 비교 현황", "모니터링 이력"],
        )
        self.assertEqual(workbook["재무데이터 입력"]["E5"].value, "(주)동남기계")
        self.assertEqual(workbook["재무데이터 입력"]["E25"].value, 59357.0)
        self.assertEqual(workbook["재무데이터 입력"]["E28"].value, 3005.0)
        self.assertEqual(workbook["재무데이터 입력"]["E32"].value, 2834.0)
        self.assertEqual(workbook["종합 평가표"]["D23"].value, calculate_health(424)["recommendation"])


class FinancialHealthSmokeTests(TestCase):
    databases = {"default"}

    def test_health_service_returns_error_for_missing_company(self):
        result = calculate_health(999999999)
        self.assertIn("error", result)

    def test_health_service_marks_cashflow_missing_when_no_amount_data(self):
        """현금흐름 금액 데이터가 없으면 is_missing=True, 총점 환산에서 제외."""
        result = calculate_health(400)
        cashflow_domain = next(domain for domain in result["domains"] if domain["name"] == "현금흐름")
        cashflow_item = cashflow_domain["items"][0]
        # 현금흐름 금액 데이터 없음 → is_missing=True, max_score=0
        self.assertTrue(cashflow_item["is_missing"])
        self.assertIsNone(cashflow_item["value"])
        self.assertEqual(cashflow_domain["max_score"], 0)
        # 총점은 현금흐름 제외 후 100점 환산
        self.assertLessEqual(result["total_score"], 100)
        self.assertGreater(result["total_score"], 0)
